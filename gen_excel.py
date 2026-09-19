#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Genera datos_congreso.xlsx con toda la información del dashboard."""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

master = json.load(open('/home/user/congreso_data.json'))
destacadas = json.load(open('/home/user/votaciones_destacadas.json'))
vots = json.load(open('/home/user/votaciones_2026.json'))  # persistente en el repo (no /tmp)
nombres_bol = json.load(open('/home/user/boletines_nombres.json'))
senado = json.load(open('/home/user/senado_presidencia.json'))  # voto por senador/a (mesa del Senado)

kpi, bancadas, dips = master['kpi'], master['bancadas'], master['diputados']

AZUL = '1D3A8F'; AZUL_CL = 'DCE6FA'; GRIS_CL = 'F2F4F9'
wb = Workbook()

def estilo_header(ws, fila=1, ncols=None):
    ncols = ncols or ws.max_column
    for c in range(1, ncols+1):
        cell = ws.cell(row=fila, column=c)
        cell.font = Font(bold=True, color='FFFFFF', size=11)
        cell.fill = PatternFill('solid', fgColor=AZUL)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

def anchos(ws, lista):
    for i, w in enumerate(lista, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

thin = Border(bottom=Side(style='thin', color='D0D5E2'))

# ============ 1. RESUMEN ============
ws = wb.active; ws.title = 'Resumen'
ws['A1'] = 'EL CONGRESO EN SIMPLE — Datos de la Cámara de Diputadas y Diputados de Chile'
ws['A1'].font = Font(bold=True, size=14, color=AZUL)
ws['A2'] = 'Fuente: miradaalcongreso.com (datos abiertos oficiales de la Cámara de Diputados y BCN)'
ws['A2'].font = Font(italic=True, color='666666')
filas = [
    ('', ''),
    ('Período legislativo', kpi['periodo']),
    ('Diputadas y diputados', 155),
    ('Sesiones de Sala contabilizadas', kpi['total_sesiones']),
    ('Fecha última actualización (fuente)', kpi['fecha_actualizacion']),
    ('Asistencia promedio a sesiones', kpi['asistencia_promedio']/100),
    ('Asistencia efectiva promedio (descuenta faltas justificadas)', kpi['asistencia_efectiva_prom']/100),
    ('Inasistencias totales', kpi['inasistencias_totales']),
    ('Inasistencias sin justificación válida', kpi['inasistencias_efectivas_totales']),
    ('', ''),
    ('Gasto total Cámara (mayo 2026, CLP)', kpi['gasto_total_camara']),
    ('Gasto promedio por diputado (mayo 2026, CLP)', kpi['gasto_promedio']),
    ('Gasto en personal de apoyo / asesores', kpi['gasto_personal_total']),
    ('Gasto operacional (traslación, telefonía, etc.)', kpi['gasto_operacional_total']),
    ('', ''),
    ('Votaciones enero–septiembre 2026', len(vots)),
    ('Aprobadas', sum(1 for v in vots if v['resultado']=='Aprobado')),
    ('Rechazadas', sum(1 for v in vots if v['resultado']=='Rechazado')),
    ('Votaciones de proyectos de ley', sum(1 for v in vots if v['tipo']=='Proyecto de Ley')),
]
for i, (a, b) in enumerate(filas, 4):
    ws.cell(row=i, column=1, value=a).font = Font(bold=bool(a))
    c = ws.cell(row=i, column=2, value=b)
    if isinstance(b, float) and b <= 1: c.number_format = '0.0%'
    if isinstance(b, int): c.number_format = '#,##0'
verif = [
    ('', ''),
    ('VERIFICACIÓN OFICIAL (17-sep-2026)', ''),
    ('Mesa de la Cámara (camara.cl)', 'Jorge Alessandri (P) · Felipe Camaño (1er VP) · Ximena Ossandón (2ª VP)'),
    ('Elección presidente Cámara', 'Alessandri 78 votos · Pamela Jiles 75 · Juan M. Valenzuela 1'),
    ('Mesa del Senado (senado.cl)', 'Paulina Núñez (P) · Iván Moreira (VP)'),
    ('Elección presidenta Senado', 'Núñez 39 votos · Alejandro Kusanovic 2 · 9 abstenciones'),
    ('Elección vicepresidente Senado', 'Moreira 37 votos a favor · 12 abstenciones'),
    ('Sesiones período actual', '69 sesiones de Sala entre el 11-mar y el 17-sep-2026 (registro oficial)'),
    ('Gastos mayo 2026', 'Chequeados contra Transparencia camara.cl: coinciden los totales de asesores y operacionales'),
]
for a, b in verif:
    ws.append([a, b])
    if a == 'VERIFICACIÓN OFICIAL (17-sep-2026)':
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12, color=AZUL)
ws.append(['', ''])
ws['A' + str(ws.max_row)] = 'Nota: este archivo es un resumen educativo generado a partir de datos públicos. Montos en pesos chilenos.'
ws['A' + str(ws.max_row)].font = Font(italic=True, size=9, color='888888')
anchos(ws, [52, 60])

# ============ 2. DIPUTADOS ============
ws = wb.create_sheet('Diputados')
hdr = ['Nombre', 'Militancia', 'Bancada', 'Asistencia real (%)', 'Asistencia efectiva (%)',
       'Inasistencias', 'Inasistencias sin justificar', 'Gasto total mayo 2026 (CLP)',
       'Gasto en asesores (CLP)', 'N° asesores', 'Gasto operacional (CLP)']
ws.append(hdr); estilo_header(ws)
for r in sorted(dips, key=lambda x: x['nombre']):
    ws.append([r['nombre'], r['militancia'], r['grupo'], r['asist']/100, r['asist_ef']/100,
               r['inas'], r['inas_ef'], r['g_tot'], r['g_per'], r['n_ases'], r['g_op']])
for row in ws.iter_rows(min_row=2):
    row[3].number_format = '0.0%'; row[4].number_format = '0.0%'
    for ci in (7, 8, 10): row[ci].number_format = '#,##0'
    for cell in row: cell.border = thin
ws.freeze_panes = 'A2'
ws.auto_filter.ref = ws.dimensions
anchos(ws, [28, 12, 20, 12, 13, 12, 15, 16, 14, 10, 14])

# ============ 3. BANCADAS ============
ws = wb.create_sheet('Bancadas')
hdr = ['Bancada', 'N° diputados', 'Asistencia promedio (%)', 'Inasistencias promedio',
       'Inasist. s/justificar promedio', 'Gasto promedio por diputado (CLP)', 'Gasto total bancada (CLP)']
ws.append(hdr); estilo_header(ws)
for b in sorted(bancadas, key=lambda x: -x['asist_prom']):
    ws.append([b['grupo'], b['n'], b['asist_prom']/100, b['inas_prom'], b['inas_ef_prom'], b['g_prom'], b['g_tot']])
for row in ws.iter_rows(min_row=2):
    row[2].number_format = '0.0%'
    for ci in (3,4,5,6): row[ci].number_format = '#,##0'
    for cell in row: cell.border = thin
ws.freeze_panes = 'A2'
anchos(ws, [24, 12, 15, 14, 17, 17, 17])

# ============ 4. VOTACIONES 2026 ============
ws = wb.create_sheet('Votaciones 2026')
ws['A1'] = 'Resumen de votaciones de la Sala — enero a septiembre de 2026'
ws['A1'].font = Font(bold=True, size=13, color=AZUL)
from collections import Counter
tipos = Counter(v['tipo'] for v in vots); res = Counter(v['resultado'] for v in vots); meses = Counter(v['fecha'][:7] for v in vots)
ws.append([]); ws.append(['Tipo de votación', 'Cantidad'])
r0 = ws.max_row
for t, n in tipos.most_common(): ws.append([t, n])
estilo_header(ws, fila=r0, ncols=2)
ws.append([]); ws.append(['Resultado', 'Cantidad'])
r0 = ws.max_row
for t, n in res.most_common(): ws.append([t, n])
estilo_header(ws, fila=r0, ncols=2)
ws.append([]); ws.append(['Mes', 'Votaciones'])
r0 = ws.max_row
for m, n in sorted(meses.items()): ws.append([m, n])
estilo_header(ws, fila=r0, ncols=2)
ws.append([]); ws.append(['Los proyectos con más votaciones del año', ''])
r0 = ws.max_row
ws.append(['Boletín', 'Nombre del proyecto', 'N° de votaciones'])
estilo_header(ws, fila=r0+1, ncols=3)
desc_count = Counter(v['desc'] for v in vots if v['desc'].startswith('Boletín'))
for b, n in desc_count.most_common(15):
    num = b.replace('Boletín N° ', '')
    if num in nombres_bol:
        ws.append([num, nombres_bol[num], n])
anchos(ws, [14, 74, 14])

# ============ 5. VOTACIONES DESTACADAS ============
ws = wb.create_sheet('Destacadas')
ws['A1'] = 'Votaciones destacadas: cómo votó cada bancada'
ws['A1'].font = Font(bold=True, size=13, color=AZUL)
for d in destacadas:
    ws.append([])
    ws.append([d['titulo'] + f" — {d['ref']}"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
    ws.append([f"Resultado oficial: {d['resultado'].upper()} · A favor {d['totales']['si']} · En contra {d['totales']['no']} · Abstención {d['totales']['abst']}"])
    ws.append(['Bancada', 'A favor', 'En contra', 'Abstención', 'Total votos'])
    estilo_header(ws, fila=ws.max_row, ncols=5)
    for g, c in sorted(d['por_bancada'].items(), key=lambda kv: -(kv[1]['si']+kv[1]['no']+kv[1]['abst'])):
        t = c['si']+c['no']+c['abst']
        if t: ws.append([g, c['si'], c['no'], c['abst'], t])
anchos(ws, [22, 10, 11, 12, 12])

# ============ 6. SENADO ============
ws = wb.create_sheet('Senado-Presidencia')
ws['A1'] = 'Elección de la Presidencia del Senado, período 2026–2030 (votación del 11 de marzo de 2026)'
ws['A1'].font = Font(bold=True, size=12, color=AZUL)
ws['A2'] = 'Resultado: Paulina Núñez electa presidenta con 39 votos · Alejandro Kusanovic 2 · Abstenciones 9 · Vicepresidencia: Iván Moreira (37 votos)'
ws['A2'].font = Font(italic=True, color='666666')
ws.append([]); ws.append(['Senador/a', 'Votó por'])
estilo_header(ws, fila=4, ncols=2)
for s, v in sorted(senado.items()):
    ws.append([s, v])
ws.freeze_panes = 'A5'
anchos(ws, [26, 22])

# ============ 7. CONSTATACIÓN OFICIAL DE ASISTENCIA ============
ws = wb.create_sheet('Asistencia oficial camara.cl')
oficial = json.load(open('/home/user/verificacion_oficial.json'))
hdr = ['Diputad@ (nombre oficial)', 'Partido', 'Sesiones asistidas (oficial)',
       'Faltas justificadas que NO descuentan', 'Faltas justificadas que SÍ descuentan',
       'Faltas sin justificación', '% asistencia oficial',
       '% asistencia dashboard (real)', '% asistencia dashboard (efectiva)', '¿Coinciden? (±1 pp)']
ws.append(hdr); estilo_header(ws)
of_by_id = {r['id']: r for r in oficial}
coinciden = 0
for d in sorted(dips, key=lambda x: x['nombre']):
    o = of_by_id.get(d['id'])
    if not o: continue
    tot = o['asiste'] + o['noafecta'] + o['rebaja'] + o['sinjust']
    pct_of = round((o['asiste'] + o['noafecta'])/tot*100, 1) if tot else None
    ok = 'Sí' if (pct_of is not None and abs(d['asist_ef'] - pct_of) <= 1.0) else 'No (licencias/permisos)'
    if pct_of is not None and abs(d['asist_ef'] - pct_of) <= 1.0: coinciden += 1
    ws.append([o['nombre'], o['partido'], o['asiste'], o['noafecta'], o['rebaja'],
               o['sinjust'], pct_of/100 if pct_of is not None else None,
               d['asist']/100, d['asist_ef']/100, ok])
for row in ws.iter_rows(min_row=2):
    for ci in (6, 7, 8):
        if row[ci].value is not None: row[ci].number_format = '0.0%'
    for cell in row: cell.border = thin
ws.append([])
ws.append([f'Conclusión: {coinciden} de {len(of_by_id)} diputad@s coinciden con el dato oficial en menos de 1 punto porcentual.'])
ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color=AZUL)
ws.append(['Las diferencias restantes se explican por licencias médicas y permisos válidos: la Cámara los descuenta del porcentaje; el tablero muestra además la cifra "real" con todas las ausencias.'])
ws.cell(row=ws.max_row, column=1).font = Font(italic=True, color='666666')
ws.append(['Fuente oficial: https://www.camara.cl/legislacion/sala_sesiones/asistencia_resumen.aspx (período 01-ene a 17-sep-2026).'])
ws.cell(row=ws.max_row, column=1).font = Font(italic=True, color='666666')
ws.freeze_panes = 'A2'
anchos(ws, [38, 22, 13, 15, 15, 13, 12, 14, 15, 18])

wb.save('/home/user/datos_congreso.xlsx')
print('OK datos_congreso.xlsx')

#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Añade DATA.presidencial.por_distrito al snapshot — sin tocar gen_elecciones.py.

Lee el MISMO ZIP oficial que ya consume la página grande
(PRELIMINARES_PRESIDENTE_DE_LA_REPUBLICA.zip, servel.cl) desde la caché de /tmp si
existe, o lo baja con la misma conexión. Cada acta de mesa trae las columnas
`distrito` + `comuna`, así que el resultado por mesa se puede agregar por los 28
distritos electorales sin pedir ningún dato nuevo.

No hay base de datos: se fusiona una clave extra en elecciones_resultados.json
(todas las claves que la página grande ya lee quedan byte a byte igual) y nada más.
"""
import datetime
import gzip
import io
import json
import os
import re as _re
import urllib.request
import zipfile

UA = 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/126 Safari/537.36'
URL = 'https://www.servel.cl/wp-content/uploads/2025/11/PRELIMINARES_PRESIDENTE_DE_LA_REPUBLICA.zip'
CACHE = '/tmp/servel_zip/presidencial.zip'
SNAP = '/home/user/elecciones_resultados.json'
COLS = ['distrito', 'comuna', 'nombre_candidato', 'votos_preliminares', 'cod_mesa', 'electores']


def descargar():
    if os.path.exists(CACHE) and os.path.getsize(CACHE) > 1000:
        return CACHE
    os.makedirs(os.path.dirname(CACHE), exist_ok=True)
    req = urllib.request.Request(URL, headers={'User-Agent': UA, 'Accept-Encoding': 'gzip'})
    with urllib.request.urlopen(req, timeout=300) as r:
        data = r.read()
        if r.headers.get('Content-Encoding') == 'gzip':
            data = gzip.GzipFile(fileobj=io.BytesIO(data)).read()
    open(CACHE, 'wb').write(data)
    print(f'  descargado presidencial: {len(data)//1024} KB')
    return CACHE


def main():
    from openpyxl import load_workbook
    z = zipfile.ZipFile(descargar())
    distritos = {}

    def dd(k):
        return distritos.setdefault(k, {'cand': {}, 'nulos': 0, 'blancos': 0,
                                        'mesas': set(), 'elec': {}})
    filas = 0
    for nom in z.namelist():
        if not nom.endswith('.xlsx'):
            continue
        wb = load_workbook(io.BytesIO(z.read(nom)), read_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        cabeza = [str(c).strip().lower() if c is not None else '' for c in next(it)]
        idx = {}
        for c in COLS:
            idx[c] = cabeza.index(c) if c in cabeza else None
        if idx['distrito'] is None or idx['votos_preliminares'] is None:
            print(f'    [salto] {nom}: no trae columnas de distrito')
            wb.close()
            continue
        for fila in it:
            filas += 1
            dis = fila[idx['distrito']]
            m = _re.search(r'(\d+)', str(dis or ''))
            if not m:
                continue
            d = dd(int(m.group(1)))
            nombre = str(fila[idx['nombre_candidato']] or '').strip()
            nu = nombre.upper()
            try:
                v = int(fila[idx['votos_preliminares']] or 0)
            except (TypeError, ValueError):
                v = 0
            mesa = str(fila[idx['cod_mesa']] or '') if idx['cod_mesa'] is not None else ''
            if mesa:
                d['mesas'].add(mesa)
                if idx['electores'] is not None:
                    try:
                        d['elec'].setdefault(mesa, int(fila[idx['electores']] or 0))
                    except (TypeError, ValueError):
                        pass
            if nu == 'VOTOS NULOS':
                d['nulos'] += v
            elif nu in ('VOTOS BLANCOS', 'VOTOS EN BLANCO'):
                d['blancos'] += v
            elif not nu.startswith(('TOTAL', 'SUMA', 'SUFRAGIOS')) and nombre:
                clave = nombre.title()
                d['cand'][clave] = d['cand'].get(clave, 0) + v
        wb.close()
        print(f'    · {nom}: acumuladas {filas} filas · {len(distritos)} distritos')

    if len(distritos) < 20:
        raise SystemExit(f'[gen_pdist] inesperado: solo {len(distritos)} distritos (aborto sin tocar el snapshot)')

    out = {}
    for k, d in sorted(distritos.items()):
        tv = sum(d['cand'].values())
        cands = sorted(d['cand'].items(), key=lambda kv: -kv[1])
        out[str(k)] = {
            'mesas': len(d['mesas']),
            'electores': sum(d['elec'].values()),
            'total_validos': tv, 'nulos': d['nulos'], 'blancos': d['blancos'],
            'candidatos': [{'nombre': n, 'votos': v, 'pct': round(v / tv * 100, 2) if tv else 0}
                           for n, v in cands[:8]],
            'ganador': cands[0][0] if cands else None,
        }

    data = json.load(open(SNAP))
    data['presidencial']['por_distrito'] = out
    data['presidencial']['fuente_por_distrito'] = {
        'archivo': URL, 'descargado': datetime.datetime.now().strftime('%Y-%m-%d %H:%M'),
        'nota': 'agregado por gen_presidencia_distritos.py a partir del mismo ZIP oficial de mesa a mesa'}
    json.dump(data, open(SNAP, 'w'), ensure_ascii=False)
    gan = {}
    for v in out.values():
        gan[v['ganador']] = gan.get(v['ganador'], 0) + 1
    print(f'OK snapshot actualizado: presidencial.por_distrito con {len(out)} distritos')
    print('   ganadores por distrito:', gan)


if __name__ == '__main__':
    main()
